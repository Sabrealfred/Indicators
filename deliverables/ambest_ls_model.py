"""Reverse-engineered AM Best life settlement securitization rating model.

Pipeline (per AM Best's published description of its approach):
  1. Eligibility gates (lives, concentration, LE providers, LE age).
  2. Mortality: base table x underwriter mortality rating, adjusted by a
     Basic Adjustment Factor, an Age-Based Adjustment Factor and a ratings
     wear-off; plus mortality improvement.
  3. Monte Carlo simulation of deaths (with a systematic LE-misestimation
     shock), insurer defaults and premium-driven lapses through the waterfall.
  4. Note default probability vs. expected maturity mapped on an idealized
     default matrix -> issue rating; stress scenarios and qualitative notching.

Parameter VALUES that AM Best does not publish in accessible form are
estimates and live in the workbook as blue/yellow inputs.

Usage:
  python ambest_ls_model.py build  Modelo.xlsx   # template + example tape
  python ambest_ls_model.py run    Modelo.xlsx   # Monte Carlo -> Resultados_MC
  python ambest_ls_model.py size   Modelo.xlsx a # Senior/Mezz split grid vs. target rating
"""
import sys
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

N_POL_MAX = 400          # tape rows reserved
LE_YEARS = 45            # horizon for life expectancy calc
AGE_MIN, AGE_MAX = 50, 120
RATINGS = ["aaa", "aa+", "aa", "aa-", "a+", "a", "a-", "bbb+", "bbb", "bbb-",
           "bb+", "bb", "bb-", "b+", "b", "b-"]
P1_INDICATIVE = [0.0008, 0.0012, 0.0020, 0.0030, 0.0040, 0.0050, 0.0065, 0.0075,
                 0.0084, 0.0123, 0.0180, 0.0260, 0.0360, 0.0500, 0.0700, 0.0950]
FSR = ["A++", "A+", "A", "A-", "B++", "B+"]
FSR_PD = [0.0001, 0.0003, 0.0006, 0.0010, 0.0025, 0.0050]

# ---- workbook layout shared by build and run --------------------------------
SUP = {  # Supuestos!C<row>
    "sr_bal": 5, "sr_cpn": 6, "mz_bal": 7, "mz_cpn": 8, "reserve": 9,
    "res_months": 10, "reinv": 11, "fix_exp": 12, "var_exp": 13,
    "remnant": 14, "T": 15, "n_sims": 16, "seed": 17, "mi": 18,
    "sys_sigma": 19, "wearoff": 20, "lgd": 21, "coi": 22, "mr_blend": 23,
    "min_lives": 24, "max_conc": 25, "max_le_age": 26, "min_prov": 27,
}
TAPE_ROW0 = 6  # Cartera_Polizas first data row
SCEN_ROW0, N_SCEN = 6, 6
RES_ROW0 = 6

ARIAL = "Arial"
BLUE = Font(name=ARIAL, size=10, color="0000FF")
BLACK = Font(name=ARIAL, size=10)
GREEN = Font(name=ARIAL, size=10, color="008000")
BOLD = Font(name=ARIAL, size=10, bold=True)
TITLE = Font(name=ARIAL, size=14, bold=True)
HDR = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F3864")
YEL = PatternFill("solid", fgColor="FFFF00")
GREY = PatternFill("solid", fgColor="F2F2F2")
MM = '$#,##0.00;($#,##0.00);-'
PCT = '0.0%;(0.0%);-'
PCT2 = '0.00%;(0.00%);-'
PCT3 = '0.000%;(0.000%);-'
NUM = '#,##0;(#,##0);-'
X = '0.00x'


def put(ws, ref, v, font=BLACK, fmt=None, fill=None):
    c = ws[ref]
    c.value = v
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def hdr(ws, row, cols, start=1):
    for i, t in enumerate(cols):
        c = ws.cell(row=row, column=start + i, value=t)
        c.font = HDR
        c.fill = HFILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


# =============================================================================
# BUILD
# =============================================================================
def example_tape(n=300, total_face=450.0, seed=7):
    rng = np.random.default_rng(seed)
    age = np.clip(np.round(rng.normal(83, 5, n)), 70, 95).astype(int)
    face = rng.lognormal(0, 0.45, n)
    face = face / face.sum() * total_face
    face = np.minimum(face, 0.025 * total_face)
    face = face / face.sum() * total_face
    mr1 = np.clip(np.round(rng.lognormal(np.log(2.0), 0.45, n) * 4) / 4, 1.0, 6.0)
    mr2 = np.clip(np.round(mr1 * rng.lognormal(0, 0.2, n) * 4) / 4, 1.0, 6.0)
    prem = np.round(0.02 + 0.0025 * (age - 70) + rng.normal(0, 0.004, n), 4)
    pf = (rng.random(n) < 0.25).astype(int)
    carriers = rng.choice(8, n, p=[.2, .18, .15, .12, .12, .1, .08, .05])
    le_age = rng.integers(3, 12, n)
    sex = rng.choice(["M", "F"], n, p=[.6, .4])
    dis = rng.choice(["Cardiovascular", "Cáncer", "Pulmonar", "Neurológica",
                      "Diabetes", "Múltiple", "Sin deterioro mayor"], n)
    return age, sex, face, carriers, prem, pf, mr1, mr2, le_age, dis


def mort_table():
    """Gompertz proxy for an ultimate VBT-style table (replace with 2015 VBT)."""
    ages = np.arange(AGE_MIN, AGE_MAX + 1)
    mu = 0.07 * np.exp(0.11 * (ages - 85))
    q = 1 - np.exp(-mu)
    q[-1] = 1.0
    return ages, np.round(q, 6)


def build(path):
    wb = Workbook()
    # ---------------- LEAME ----------------
    ws = wb.active
    ws.title = "LEAME"
    txt = [
        ("Ingeniería inversa – Metodología AM Best 'Life Settlement Securitization' (SPV USD 100 mm)", TITLE),
        ("Réplica independiente del proceso descrito públicamente por AM Best. NO es el modelo propietario ni una calificación de AM Best.", BOLD),
        ("", None),
        ("Proceso replicado (en el orden que AM Best lo describe)", BOLD),
        ("1. Filtros de elegibilidad: ~300 vidas, ≤3.33% del face por vida, ≥2 proveedores de LE, LE ≤12 meses, cartera adquirida antes de calificar.", None),
        ("2. Mortalidad: tabla base (2015 VBT en la revisión 2026) × calificación de mortalidad del underwriter, ajustada por", None),
        ("   Basic Adjustment Factor (nivel de la calificación y si la póliza fue premium-financed), Age-Based Adjustment Factor y Ratings Wear-Off.", None),
        ("3. Simulación Monte Carlo de muertes (con choque sistemático por error de estimación de LE), default de aseguradoras y caducidad por falta de liquidez.", None),
        ("4. Probabilidad de default de cada nota vs. su vencimiento esperado → Matriz Idealizada de Default (BILSM, Exhibit E.2) → calificación.", None),
        ("5. Escenarios de estrés y factores cualitativos → ajuste en notches para el comité.", None),
        ("", None),
        ("Cómo usarlo", BOLD),
        ("a) Pegue su cinta póliza por póliza en 'Cartera_Polizas' (columnas azules). Las columnas negras se calculan solas.", None),
        ("b) Reemplace 'Tabla_Mortalidad' por la 2015 VBT real (q_x por edad) y 'Matriz_Default' por el Exhibit E.2 de BILSM.", None),
        ("c) Ajuste 'Factores' (BAF, AAF, PD de aseguradoras) y 'Supuestos'.", None),
        ("d) Ejecute:  python ambest_ls_model.py run Modelo.xlsx   → llena 'Resultados_MC'. Excel recalcula las calificaciones al abrir.", None),
        ("", None),
        ("Leyenda: azul = input · negro = fórmula · verde = vínculo a otra hoja · amarillo = supuesto clave (estimación a calibrar).", None),
        ("Los valores de BAF, AAF, wear-off, sigma sistemática y la matriz de default son ESTIMACIONES: AM Best describe estos factores pero sus tablas no fueron accesibles.", None),
    ]
    for i, (t, f) in enumerate(txt, 1):
        ws.cell(row=i, column=1, value=t).font = f or BLACK
    ws.column_dimensions["A"].width = 150

    # ---------------- Supuestos ----------------
    S = wb.create_sheet("Supuestos")
    put(S, "A1", "Supuestos", TITLE)
    hdr(S, 3, ["", "Parámetro", "Valor", "Unidad", "Fuente / estatus"])
    rows = [
        ("sr_bal", "Notas Senior", 75, "USD mm", MM, "Estructura ilustrativa"),
        ("sr_cpn", "Cupón Senior", 0.065, "% anual", PCT2, "Estructura ilustrativa"),
        ("mz_bal", "Notas Mezzanine", 25, "USD mm", MM, "Estructura ilustrativa"),
        ("mz_cpn", "Cupón Mezzanine", 0.095, "% anual", PCT2, "Estructura ilustrativa"),
        ("reserve", "Reserva de liquidez inicial (equity)", 20, "USD mm", MM, "Estructura ilustrativa"),
        ("res_months", "Reserva objetivo (meses de primas + gastos)", 12, "meses", NUM, "Práctica de mercado"),
        ("reinv", "Tasa de reinversión de caja", 0.03, "% anual", PCT2, "Supuesto"),
        ("fix_exp", "Gastos fijos (servicer, trustee, tracking)", 0.75, "USD mm/año", MM, "Supuesto"),
        ("var_exp", "Gasto variable sobre face en vigor", 0.0005, "% anual", PCT3, "Supuesto"),
        ("remnant", "Venta de cartera remanente al vencimiento legal", 0.20, "% face", PCT, "Supuesto (AM Best suele dar poco crédito)"),
        ("T", "Vencimiento legal (años, máx. 40)", 25, "años", NUM, "Estructura ilustrativa"),
        ("n_sims", "Número de simulaciones Monte Carlo", 5000, "#", NUM, "Motor (AM Best: Monte Carlo)"),
        ("seed", "Semilla aleatoria", 2026, "#", NUM, "Reproducibilidad"),
        ("mi", "Mejora anual de mortalidad", 0.005, "% anual", PCT2, "Estimación (reduce q = extiende LE)"),
        ("sys_sigma", "Sigma del error sistemático de LE (lognormal)", 0.20, "σ", PCT, "Estimación – AM Best cita 'systematic misestimation'"),
        ("wearoff", "Años de desvanecimiento (wear-off) del deterioro", 15, "años", NUM, "Estimación – AM Best: 'Ratings Wear-Off Adjustment'"),
        ("lgd", "Pérdida dado default de aseguradora", 0.50, "% beneficio", PCT, "Estimación"),
        ("coi", "Aumento anual de primas (COI)", 0.01, "% anual", PCT2, "Estimación"),
        ("mr_blend", "Combinación de las 2 LE: 1 = promedio, 2 = conservadora (LE más larga)", 1, "1/2", NUM, "Estimación"),
        ("min_lives", "Mínimo de vidas sin estrés adicional", 300, "vidas", NUM, "AM Best (publicado)"),
        ("max_conc", "Concentración máxima por vida", 0.0333, "% face", PCT2, "AM Best (publicado)"),
        ("max_le_age", "Antigüedad máxima de LE", 12, "meses", NUM, "AM Best (publicado)"),
        ("min_prov", "Mínimo de proveedores de LE", 2, "#", NUM, "AM Best (publicado)"),
    ]
    for key, lbl, v, u, fmt, src in rows:
        r = SUP[key]
        put(S, f"B{r}", lbl)
        put(S, f"C{r}", v, BLUE, fmt, YEL if key in ("sys_sigma", "wearoff", "mi", "lgd", "coi") else None)
        put(S, f"D{r}", u)
        put(S, f"E{r}", src)
    for col, w in zip("ABCDE", [2, 62, 12, 12, 60]):
        S.column_dimensions[col].width = w

    # ---------------- Factores ----------------
    Fa = wb.create_sheet("Factores")
    put(Fa, "A1", "Factores de ajuste de mortalidad y riesgo de aseguradoras", TITLE)
    put(Fa, "A2", "Estructura descrita por AM Best; los valores son estimaciones a calibrar (factor < 1 reduce la mortalidad = extiende la LE).")
    put(Fa, "A4", "Basic Adjustment Factor (BAF) – por nivel de calificación de mortalidad y premium financing", BOLD)
    hdr(Fa, 5, ["MR desde (≥)", "BAF no PF", "BAF premium-financed"])
    baf = [(0, 0.95, 0.90), (1.5, 0.88, 0.82), (3.0, 0.80, 0.74), (5.0, 0.72, 0.66)]
    for i, (a, b, c) in enumerate(baf):
        r = 6 + i
        put(Fa, f"A{r}", a, BLUE, X)
        put(Fa, f"B{r}", b, BLUE, X, YEL)
        put(Fa, f"C{r}", c, BLUE, X, YEL)
    put(Fa, "A11", "Age-Based Adjustment Factor (AAF)", BOLD)
    hdr(Fa, 12, ["Edad desde (≥)", "AAF"])
    aaf = [(0, 0.90), (75, 0.95), (85, 1.00), (90, 1.00)]
    for i, (a, b) in enumerate(aaf):
        r = 13 + i
        put(Fa, f"A{r}", a, BLUE, NUM)
        put(Fa, f"B{r}", b, BLUE, X, YEL)
    put(Fa, "A18", "PD anual de aseguradoras por FSR", BOLD)
    hdr(Fa, 19, ["FSR", "PD anual"])
    for i, (f, p) in enumerate(zip(FSR, FSR_PD)):
        r = 20 + i
        put(Fa, f"A{r}", f, BLUE)
        put(Fa, f"B{r}", p, BLUE, PCT3, YEL)
    put(Fa, "A27", "Aseguradoras de la cartera", BOLD)
    hdr(Fa, 28, ["ID", "Nombre", "FSR"])
    car_fsr = ["A+", "A+", "A", "A", "A-", "A", "A+", "B++"]
    for i in range(8):
        r = 29 + i
        put(Fa, f"A{r}", i + 1, BLUE, NUM)
        put(Fa, f"B{r}", f"Aseguradora {i + 1}", BLUE)
        put(Fa, f"C{r}", car_fsr[i], BLUE)
    for col, w in zip("ABC", [16, 18, 22]):
        Fa.column_dimensions[col].width = w

    # ---------------- Tabla_Mortalidad ----------------
    M = wb.create_sheet("Tabla_Mortalidad")
    put(M, "A1", "Tabla base de mortalidad q_x (proxy Gompertz – REEMPLAZAR por 2015 VBT)", TITLE)
    put(M, "A2", "Proxy: μ(x) = 0.07·e^{0.11(x−85)}, q = 1−e^{−μ}; q(120) = 1. La revisión 2026 de AM Best usa 2015 VBT.")
    hdr(M, 4, ["Edad", "q_x"])
    ages, q = mort_table()
    for i, (a, v) in enumerate(zip(ages, q)):
        put(M, f"A{5 + i}", int(a), BLACK, NUM)
        put(M, f"B{5 + i}", float(v), BLUE, '0.000000')
    QROW0, QROW1 = 5, 5 + len(ages) - 1

    # ---------------- Cartera_Polizas ----------------
    C = wb.create_sheet("Cartera_Polizas")
    put(C, "A1", "Cinta de datos póliza por póliza (ejemplo sintético de 300 vidas – reemplazar)", TITLE)
    put(C, "A2", "Columnas A–L = input. M–R = fórmulas (ajustes de AM Best replicados). Deje filas vacías al final; el motor lee hasta la primera fila sin ID.")
    cols = ["ID", "Edad", "Sexo", "Face (USD mm)", "Aseguradora ID", "FSR", "Prima anual (% face)", "Premium financed (1/0)",
            "MR proveedor 1", "MR proveedor 2", "Antigüedad LE (meses)", "Enfermedad principal",
            "MR combinado", "BAF", "AAF", "MR ajustado", "LE ajustada (años)", "% del face"]
    hdr(C, 5, cols)
    age, sex, face, car, prem, pf, mr1, mr2, le_age, dis = example_tape()
    n = len(age)
    for i in range(n):
        r = TAPE_ROW0 + i
        vals = [i + 1, int(age[i]), sex[i], float(round(face[i], 4)), int(car[i]) + 1, None, float(prem[i]), int(pf[i]),
                float(mr1[i]), float(mr2[i]), int(le_age[i]), dis[i]]
        fmts = [NUM, NUM, None, MM, NUM, None, PCT2, NUM, X, X, NUM, None]
        for j, (v, f) in enumerate(zip(vals, fmts)):
            if j == 5:
                put(C, f"F{r}", f"=INDEX(Factores!$C$29:$C$36,MATCH(E{r},Factores!$A$29:$A$36,0))", GREEN)
            else:
                put(C, f"{L(j + 1)}{r}", v, BLUE, f)
        put(C, f"M{r}", f"=IF(Supuestos!$C${SUP['mr_blend']}=1,AVERAGE(I{r}:J{r}),MIN(I{r}:J{r}))", BLACK, X)
        put(C, f"N{r}", f"=INDEX(IF(H{r}=1,Factores!$C$6:$C$9,Factores!$B$6:$B$9),MATCH(M{r},Factores!$A$6:$A$9,1))", BLACK, X)
        put(C, f"O{r}", f"=INDEX(Factores!$B$13:$B$16,MATCH(B{r},Factores!$A$13:$A$16,1))", BLACK, X)
        put(C, f"P{r}", f"=M{r}*N{r}*O{r}", BLACK, X)
        put(C, f"Q{r}", f"=LE!B{4 + i}", GREEN, '0.00')
        put(C, f"R{r}", f"=D{r}/$D$3", BLACK, PCT2)
    last = TAPE_ROW0 + n - 1
    put(C, "C3", "Face total", BOLD)
    put(C, "D3", f"=SUM(D{TAPE_ROW0}:D{TAPE_ROW0 + N_POL_MAX - 1})", BLACK, MM)
    put(C, "F3", "Vidas", BOLD)
    put(C, "G3", f"=COUNT(A{TAPE_ROW0}:A{TAPE_ROW0 + N_POL_MAX - 1})", BLACK, NUM)
    put(C, "H3", "LE ponderada", BOLD)
    put(C, "I3", f"=SUMPRODUCT(Q{TAPE_ROW0}:Q{last},D{TAPE_ROW0}:D{last})/D3", BLACK, '0.00')
    put(C, "J3", "Mayor póliza %", BOLD)
    put(C, "K3", f"=MAX(R{TAPE_ROW0}:R{last})", BLACK, PCT2)
    put(C, "L3", "LE más antigua (m)", BOLD)
    put(C, "M3", f"=MAX(K{TAPE_ROW0}:K{last})", BLACK, NUM)
    put(C, "N3", "MR ajust. prom.", BOLD)
    put(C, "O3", f"=SUMPRODUCT(P{TAPE_ROW0}:P{last},D{TAPE_ROW0}:D{last})/D3", BLACK, X)
    widths = [6, 6, 6, 12, 10, 6, 10, 10, 10, 10, 10, 20, 11, 8, 8, 10, 10, 9]
    for j, w in enumerate(widths):
        C.column_dimensions[L(j + 1)].width = w
    C.freeze_panes = "B6"

    # ---------------- LE helper ----------------
    E = wb.create_sheet("LE")
    put(E, "A1", "Supervivencia ajustada por póliza: q_t = MIN(1, q_base(edad+t−1) × [1+(MRaj−1)·máx(0,1−(t−1)/W)] × (1−MI)^(t−1));  LE = Σ S_t + 0.5", BOLD)
    hdr(E, 3, ["ID", "LE (años)"] + [str(t) for t in range(1, LE_YEARS + 1)])
    w = f"Supuestos!$C${SUP['wearoff']}"
    mi = f"Supuestos!$C${SUP['mi']}"
    for i in range(n):
        r = 4 + i
        cr = TAPE_ROW0 + i
        put(E, f"A{r}", f"=Cartera_Polizas!A{cr}", GREEN)
        put(E, f"B{r}", f"=SUM(C{r}:{L(2 + LE_YEARS)}{r})+0.5", BLACK, '0.00')
        for t in range(1, LE_YEARS + 1):
            col = L(2 + t)
            prev = L(1 + t)
            qb = f"INDEX(Tabla_Mortalidad!$B${QROW0}:$B${QROW1},MIN({AGE_MAX - AGE_MIN + 1},Cartera_Polizas!$B${cr}+{t - 1}-{AGE_MIN - 1}))"
            mult = f"(1+(Cartera_Polizas!$P${cr}-1)*MAX(0,1-{t - 1}/{w}))"
            q_ = f"MIN(1,{qb}*{mult}*(1-{mi})^{t - 1})"
            put(E, f"{col}{r}", f"=1-{q_}" if t == 1 else f"={prev}{r}*(1-{q_})", BLACK, '0.000')

    # ---------------- Matriz_Default ----------------
    D = wb.create_sheet("Matriz_Default")
    put(D, "A1", "Matriz idealizada de default por emisión (INDICATIVA – reemplazar con BILSM Exhibit E.2)", TITLE)
    put(D, "A2", "Acumulada(t) = 1 − (1 − p1)^t. p1 parcialmente referenciado (aaa 1a ≈ 0.08%, bbb 1a ≈ 0.84%); resto interpolado por Claude.")
    hdr(D, 4, ["Calificación", "p1 (1 año)"] + [str(t) for t in range(1, 31)])
    for i, (rt, p) in enumerate(zip(RATINGS, P1_INDICATIVE)):
        r = 5 + i
        put(D, f"A{r}", rt, BOLD)
        put(D, f"B{r}", p, BLUE, PCT2, YEL)
        for t in range(1, 31):
            put(D, f"{L(2 + t)}{r}", f"=1-(1-$B{r})^{t}", BLACK, PCT2)
    put(D, "A22", "ccc o menor", BOLD)
    D.column_dimensions["A"].width = 14

    # ---------------- Escenarios ----------------
    Sc = wb.create_sheet("Escenarios")
    put(Sc, "A1", "Escenarios Monte Carlo (base + estrés para comité)", TITLE)
    put(Sc, "A2", "Multiplicador q < 1 = LE más larga. Sigma vacía = usa Supuestos. 'Default aseguradora mayor' fuerza el default de la aseguradora con más face en el año 1.")
    hdr(Sc, 5, ["Escenario", "Multiplicador q", "COI adicional", "Sigma sistemática", "Default aseguradora mayor (1/0)", "Reserva inicial × "])
    scen = [("Base", 1.0, 0.0, None, 0, 1.0), ("Extensión LE (q × 0.85)", 0.85, 0.0, None, 0, 1.0),
            ("Extensión severa (q × 0.75)", 0.75, 0.0, None, 0, 1.0), ("Primas +1.5% anual adicional", 1.0, 0.015, None, 0, 1.0),
            ("Default aseguradora mayor", 1.0, 0.0, None, 1, 1.0), ("Liquidez −50%", 1.0, 0.0, None, 0, 0.5)]
    for i, s in enumerate(scen):
        r = SCEN_ROW0 + i
        for j, (v, f) in enumerate(zip(s, [None, X, PCT2, PCT, NUM, X])):
            put(Sc, f"{L(j + 1)}{r}", v, BLUE if j else BOLD, f)
    for col, w_ in zip("ABCDEF", [32, 14, 12, 14, 18, 14]):
        Sc.column_dimensions[col].width = w_

    # ---------------- Ponderacion ----------------
    P = wb.create_sheet("Ponderacion")
    put(P, "A1", "Factores cualitativos y chequeos de elegibilidad", TITLE)
    put(P, "A2", "Pesos = estimación de Claude (AM Best no publica una ponderación). El núcleo de la calificación es la PD de Monte Carlo.")
    hdr(P, 4, ["#", "Factor", "Peso", "Puntaje (1–5)", "Ponderado"])
    facts = [("Calidad de LE / underwriting médico", 0.25, 4), ("Diversificación de la cartera", 0.20, 4),
             ("Liquidez para primas", 0.20, 3), ("Calidad crediticia de aseguradoras", 0.10, 4),
             ("Riesgo legal de pólizas (STOLI, interés asegurable)", 0.10, 3), ("Estructura legal del SPV", 0.10, 4),
             ("Contrapartes operativas", 0.05, 4)]
    for i, (f, w_, s) in enumerate(facts):
        r = 5 + i
        put(P, f"A{r}", i + 1)
        put(P, f"B{r}", f)
        put(P, f"C{r}", w_, BLUE, PCT, YEL)
        put(P, f"D{r}", s, BLUE, NUM, YEL)
        put(P, f"E{r}", f"=C{r}*D{r}", BLACK, '0.00')
    put(P, "B12", "Total", BOLD)
    put(P, "C12", "=SUM(C5:C11)", BLACK, PCT)
    put(P, "E12", "=SUM(E5:E11)", BLACK, '0.00')
    hdr(P, 14, ["", "Puntaje ≥", "Notches"])
    for i, (th, nn) in enumerate([(0, -3), (2.5, -2), (3.5, -1), (4.5, 0)]):
        put(P, f"B{15 + i}", th, BLUE, '0.0')
        put(P, f"C{15 + i}", nn, BLUE, NUM)
    put(P, "B20", "Notches cualitativos", BOLD)
    put(P, "C20", "=INDEX(C15:C18,MATCH(E12,B15:B18,1))", BLACK, NUM)
    hdr(P, 22, ["", "Chequeo de elegibilidad (AM Best)", "Valor", "Límite", "Resultado"])
    chk = [("# de vidas", "=Cartera_Polizas!G3", f"=Supuestos!C{SUP['min_lives']}", '=IF(C23>=D23,"OK","ESTRÉS ADICIONAL")', NUM),
           ("Mayor póliza / face", "=Cartera_Polizas!K3", f"=Supuestos!C{SUP['max_conc']}", '=IF(C24<=D24,"OK","EXCEDE")', PCT2),
           ("LE más antigua (meses)", "=Cartera_Polizas!M3", f"=Supuestos!C{SUP['max_le_age']}", '=IF(C25<=D25,"OK","ACTUALIZAR LE")', NUM),
           ("# proveedores LE (columnas I y J llenas)", f"=IF(COUNT(Cartera_Polizas!J{TAPE_ROW0}:J{last})=Cartera_Polizas!G3,2,1)", f"=Supuestos!C{SUP['min_prov']}", '=IF(C26>=D26,"OK","INSUFICIENTE")', NUM)]
    for i, (l_, v_, lim, res, f) in enumerate(chk):
        r = 23 + i
        put(P, f"B{r}", l_)
        put(P, f"C{r}", v_, GREEN, f)
        put(P, f"D{r}", lim, GREEN, f)
        put(P, f"E{r}", res)
    put(P, "B27", "Notches por chequeos fallidos (−1 c/u)")
    put(P, "C27", '=-COUNTIF(E23:E26,"<>OK")', BLACK, NUM)
    put(P, "B28", "Total notches", BOLD)
    put(P, "C28", "=C20+C27", BLACK, NUM)
    for col, w_ in zip("ABCDE", [4, 48, 12, 12, 18]):
        P.column_dimensions[col].width = w_

    # ---------------- Resultados_MC ----------------
    R = wb.create_sheet("Resultados_MC", 1)
    put(R, "A1", "Resultados Monte Carlo y calificación indicada (réplica – no es una calificación de AM Best)", TITLE)
    put(R, "A2", "Valores en columnas B–K los escribe el motor (python ambest_ls_model.py run). Columnas L–Q son fórmulas contra 'Matriz_Default'.")
    put(R, "A3", "Estado:")
    put(R, "B3", "Sin correr – ejecute el motor", BOLD)
    hdr(R, 5, ["Escenario", "PD Senior", "Pérdida esperada Senior", "Vencim. esperado Senior (años)", "PD Mezz", "Pérdida esperada Mezz",
               "Vencim. esperado Mezz (años)", "P(caducidad con notas vigentes)", "LE realizada prom. (años)", "Muertes cobradas prom. (USD mm)", "Sims",
               "Senior – índice", "Senior – calificación MC", "Mezz – índice", "Mezz – calificación MC", "Senior – final (con notches)", "Mezz – final (con notches)"])
    MAT = "Matriz_Default!$C$5:$AF$20"
    for i in range(N_SCEN):
        r = RES_ROW0 + i
        put(R, f"A{r}", f"=Escenarios!A{SCEN_ROW0 + i}", GREEN)
        for col, f in zip("BCDEFGHIJK", [PCT3, PCT3, '0.0', PCT3, PCT3, '0.0', PCT2, '0.00', MM, NUM]):
            R[f"{col}{r}"].number_format = f
            R[f"{col}{r}"].font = BLUE
        for pdc, tc, ic, rc in (("B", "D", "L", "M"), ("E", "G", "N", "O")):
            ten = f"MIN(30,MAX(1,ROUND({tc}{r},0)))"
            own = f'COUNTIF(INDEX({MAT},0,{ten}),"<"&{pdc}{r})+1'
            if ic == "N":  # a subordinate tranche cannot rate above the Senior
                own = f"MAX(L{r},{own})"
            put(R, f"{ic}{r}", f'=IF({pdc}{r}="","",{own})', BLACK, NUM)
            put(R, f"{rc}{r}", f'=IF({ic}{r}="","",INDEX(Matriz_Default!$A$5:$A$21,MIN(17,{ic}{r})))', BLACK)
        put(R, f"P{r}", f'=IF(L{r}="","",INDEX(Matriz_Default!$A$5:$A$21,MIN(17,L{r}-Ponderacion!$C$28)))', BLACK)
        put(R, f"Q{r}", f'=IF(N{r}="","",INDEX(Matriz_Default!$A$5:$A$21,MIN(17,N{r}-Ponderacion!$C$28)))', BLACK)
    put(Mat_fix := D, "A21", "ccc o menor", BOLD)
    put(D, "A22", None)
    fr = RES_ROW0 + N_SCEN + 1
    put(R, f"A{fr}", "CALIFICACIÓN INDICADA (caso base MC − notches cualitativos y de elegibilidad)", BOLD)
    put(R, f"A{fr + 1}", "Senior")
    put(R, f"B{fr + 1}", f"=P{RES_ROW0}", BOLD, None, YEL)
    put(R, f"A{fr + 2}", "Mezzanine")
    put(R, f"B{fr + 2}", f"=Q{RES_ROW0}", BOLD, None, YEL)
    put(R, f"A{fr + 3}", "Notches aplicados")
    put(R, f"B{fr + 3}", "=Ponderacion!C28", GREEN, NUM)
    put(R, f"A{fr + 5}", "Conciliación: LE ponderada Excel vs. motor")
    put(R, f"B{fr + 5}", "=Cartera_Polizas!I3", GREEN, '0.00')
    put(R, f"C{fr + 5}", None, BLUE, '0.00')
    put(R, f"D{fr + 5}", f'=IF(C{fr + 5}="","",IF(ABS(B{fr + 5}-C{fr + 5})<0.05,"OK","REVISAR"))')
    widths = [30, 11, 12, 12, 11, 12, 12, 12, 11, 13, 8, 9, 12, 9, 12, 13, 13]
    for j, w_ in enumerate(widths):
        R.column_dimensions[L(j + 1)].width = w_

    # ---------------- Fuentes ----------------
    Q = wb.create_sheet("Fuentes")
    put(Q, "A1", "Fuentes y estatus", TITLE)
    hdr(Q, 3, ["Elemento", "Estatus", "Fuente"])
    src = [
        ("Monte Carlo para PD de los valores", "Descrito por AM Best", "AM Best press release – https://news.ambest.com/PR/PressContent.aspx?altsrc=0&refnum=9726"),
        ("Basic Adjustment Factor, Age-Based Adjustment Factor, Ratings Wear-Off", "Descrito por AM Best; valores ESTIMADOS", "AM Best, Life Settlement Securitization – https://www3.ambest.com/ambv/ratingmethodology/OpenPDF.aspx?rc=197705"),
        ("Tabla VBT ajustada (2008 → 2015 VBT en revisión 2026)", "Descrito por AM Best; tabla aquí es PROXY", "BusinessWire 8-sep-2026 – https://secure.businesswire.com/news/home/20260908651291/en/"),
        ("PD vs. vencimiento esperado → Idealized Issue Default Matrix", "Descrito por AM Best; matriz INDICATIVA", "BILSM – https://www3.ambest.com/ambv/ratingmethodology/openpdf.aspx?ubcr=1&ri=1124 ; https://www3.ambest.com/ambv/ratingmethodology/OpenPDF.aspx?rc=197709"),
        ("≥300 vidas, ≤3.33% por vida, ≥2 LE, ≤12 meses, cartera adquirida (≤18 meses)", "Publicado por AM Best", "Life Settlement Securitization (ídem)"),
        ("Etapas: evaluación preliminar → indicativa → calificación de largo plazo", "Publicado por AM Best", "Ídem"),
        ("Sigma sistemática, COI, LGD, mejora de mortalidad, pesos cualitativos", "ESTIMACIÓN de Claude", "Calibrar con el criterio completo"),
    ]
    for i, row in enumerate(src):
        for j, v in enumerate(row):
            put(Q, f"{L(j + 1)}{4 + i}", v)
    for col, w_ in zip("ABC", [70, 38, 120]):
        Q.column_dimensions[col].width = w_

    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


# =============================================================================
# RUN (Monte Carlo engine)
# =============================================================================
def _lookup_band(x, bounds, values):
    idx = np.searchsorted(np.asarray(bounds), x, side="right") - 1
    return np.asarray(values)[np.clip(idx, 0, len(values) - 1)]


def read_inputs(path):
    wb = load_workbook(path)
    S, Fa, M, C, Sc = (wb[n] for n in ("Supuestos", "Factores", "Tabla_Mortalidad", "Cartera_Polizas", "Escenarios"))
    p = {k: S[f"C{r}"].value for k, r in SUP.items()}
    baf = [[Fa[f"{c}{r}"].value for r in range(6, 10)] for c in "ABC"]
    aaf = [[Fa[f"{c}{r}"].value for r in range(13, 17)] for c in "AB"]
    fsr_pd = {Fa[f"A{r}"].value: Fa[f"B{r}"].value for r in range(20, 26)}
    car_fsr = {Fa[f"A{r}"].value: Fa[f"C{r}"].value for r in range(29, 37) if Fa[f"A{r}"].value is not None}
    qtab = {}
    r = 5
    while M[f"A{r}"].value is not None:
        qtab[int(M[f"A{r}"].value)] = float(M[f"B{r}"].value)
        r += 1
    rows = []
    r = TAPE_ROW0
    while C[f"A{r}"].value not in (None, ""):
        rows.append([C[f"{c}{r}"].value for c in "ABCDEFGHIJK"])
        r += 1
    scen = []
    for i in range(N_SCEN):
        rr = SCEN_ROW0 + i
        if Sc[f"A{rr}"].value:
            scen.append(dict(name=Sc[f"A{rr}"].value, qmult=Sc[f"B{rr}"].value or 1.0, coi_add=Sc[f"C{rr}"].value or 0.0,
                             sigma=Sc[f"D{rr}"].value, big_carrier=int(Sc[f"E{rr}"].value or 0), res_mult=Sc[f"F{rr}"].value or 1.0))
    return wb, p, baf, aaf, fsr_pd, car_fsr, qtab, rows, scen


def policy_q(p, baf, aaf, qtab, rows, horizon):
    age = np.array([r[1] for r in rows], float)
    pf = np.array([r[7] for r in rows], int)
    mr1 = np.array([r[8] for r in rows], float)
    mr2 = np.array([r[9] if r[9] is not None else r[8] for r in rows], float)
    mr = (mr1 + mr2) / 2 if int(p["mr_blend"]) == 1 else np.minimum(mr1, mr2)
    b = np.where(pf == 1, _lookup_band(mr, baf[0], baf[2]), _lookup_band(mr, baf[0], baf[1]))
    a = _lookup_band(age, aaf[0], aaf[1])
    mr_adj = mr * b * a
    t = np.arange(horizon)
    ages = np.minimum(age[:, None] + t[None, :], AGE_MAX).astype(int)
    qb = np.vectorize(lambda x: qtab.get(int(x), 1.0))(ages)
    wear = np.maximum(0, 1 - t / float(p["wearoff"]))
    mult = 1 + (mr_adj[:, None] - 1) * wear[None, :]
    q = np.minimum(1, qb * mult * (1 - p["mi"]) ** t[None, :])
    return q


def simulate(p, q, rows, fsr_pd, car_fsr, sc):
    T = int(p["T"])
    n_s = int(p["n_sims"])
    rng = np.random.default_rng(int(p["seed"]))
    face = np.array([r[3] for r in rows], float)
    prem_rate = np.array([r[6] for r in rows], float)
    car = np.array([r[4] for r in rows], int)
    n = len(face)
    sigma = p["sys_sigma"] if sc["sigma"] is None else sc["sigma"]
    eps = np.exp(sigma * rng.standard_normal(n_s) - sigma ** 2 / 2) * sc["qmult"]
    # insurer default years
    cars = sorted(set(car.tolist()))
    pd_c = np.array([fsr_pd.get(car_fsr.get(c, "A"), 0.001) for c in cars])
    u = rng.random((n_s, len(cars)))
    dflt_year = np.ceil(np.log(u) / np.log(1 - pd_c[None, :]))
    if sc["big_carrier"]:
        big = cars[int(np.argmax([face[car == c].sum() for c in cars]))]
        dflt_year[:, cars.index(big)] = 1
    cidx = np.searchsorted(cars, car)
    coi = p["coi"] + sc["coi_add"]

    alive = np.ones((n_s, n), bool)
    k = np.ones(n_s)                      # in-force fraction (pro-rata lapse)
    cash = np.full(n_s, p["reserve"] * sc["res_mult"])
    sr = np.full(n_s, float(p["sr_bal"]))
    mz = np.full(n_s, float(p["mz_bal"]))
    sr_int_short = np.zeros(n_s, bool)
    lapsed = np.zeros(n_s, bool)
    sr_wal = np.zeros(n_s)
    mz_wal = np.zeros(n_s)
    db_tot = np.zeros(n_s)
    death_t = np.full((n_s, n), np.nan)
    for t in range(T):
        # deaths and (monthly) premiums both occur during the year: net them annually
        face_inforce = k * (alive * face).sum(1)
        prem_due = k * (alive * (face * prem_rate)).sum(1) * (1 + coi) ** t
        dies = alive & (rng.random((n_s, n)) < np.minimum(1, q[:, t][None, :] * eps[:, None]))
        death_t[dies] = t + 0.5
        recov = np.where(dflt_year[:, cidx] <= t + 1, 1 - p["lgd"], 1.0)
        db = k * (dies * face * recov).sum(1)
        income = np.maximum(cash, 0) * p["reinv"]
        avail = cash + db + income
        paid = np.minimum(np.maximum(avail, 0), prem_due)
        short = prem_due > paid * (1 + 1e-9) + 1e-9
        frac = np.where(prem_due > 0, paid / np.where(prem_due > 0, prem_due, 1), 1)
        lapsed |= short & ((sr > 1e-9) | (mz > 1e-9))
        db_tot += db
        alive &= ~dies
        k = k * np.where(short, frac, 1)       # unpaid share of the pool lapses
        exp_ = p["fix_exp"] + p["var_exp"] * face_inforce
        cash = avail - paid - exp_
        if t == T - 1:
            cash = cash + k * (alive * face).sum(1) * p["remnant"]
        # waterfall
        sr_due = sr * p["sr_cpn"]
        sr_paid = np.clip(cash, 0, sr_due)
        sr_int_short |= (sr > 1e-9) & (sr_paid < sr_due - 1e-9)
        cash -= sr_paid
        sr += sr_due - sr_paid
        mz_due = mz * p["mz_cpn"]
        mz_paid = np.clip(cash, 0, mz_due)
        cash -= mz_paid
        mz += mz_due - mz_paid
        nxt_prem = k * (alive * (face * prem_rate)).sum(1) * (1 + coi) ** (t + 1)
        nxt_exp = p["fix_exp"] + p["var_exp"] * k * (alive * face).sum(1)
        target = 0 if t == T - 1 else (nxt_prem + nxt_exp) * p["res_months"] / 12
        sr_pr = np.clip(cash - target, 0, sr)
        cash -= sr_pr
        sr -= sr_pr
        mz_pr = np.clip(cash - target, 0, mz)
        cash -= mz_pr
        mz -= mz_pr
        sr_wal += sr_pr * (t + 1)
        mz_wal += mz_pr * (t + 1)
        done = (sr < 1e-9) & (mz < 1e-9)
        cash = np.where(done, np.minimum(cash, np.maximum(target, 0)), cash)  # release excess to equity
    sr_def = sr_int_short | (sr > 1e-6)
    mz_def = mz > 1e-6
    sr_loss = np.minimum(1, sr / p["sr_bal"])
    mz_loss = np.minimum(1, mz / p["mz_bal"])
    sr_mat = np.where(sr_def, T, sr_wal / p["sr_bal"])
    mz_mat = np.where(mz_def, T, mz_wal / p["mz_bal"])
    realized_le = np.nanmean(np.where(np.isnan(death_t), T, death_t))
    return dict(pd_sr=sr_def.mean(), el_sr=sr_loss.mean(), mat_sr=sr_mat.mean(),
                pd_mz=mz_def.mean(), el_mz=mz_loss.mean(), mat_mz=mz_mat.mean(),
                p_lapse=lapsed.mean(), le=realized_le, db=db_tot.mean(), sims=n_s)


def run(path):
    wb, p, baf, aaf, fsr_pd, car_fsr, qtab, rows, scen = read_inputs(path)
    T = int(p["T"])
    q = policy_q(p, baf, aaf, qtab, rows, max(T, LE_YEARS))
    face = np.array([r[3] for r in rows], float)
    surv = np.cumprod(1 - q[:, :LE_YEARS], axis=1)
    le = surv.sum(1) + 0.5
    le_w = float((le * face).sum() / face.sum())
    R = wb["Resultados_MC"]
    for i, sc in enumerate(scen):
        res = simulate(p, q[:, :T], rows, fsr_pd, car_fsr, sc)
        r = RES_ROW0 + i
        for col, key in zip("BCDEFGHIJK", ["pd_sr", "el_sr", "mat_sr", "pd_mz", "el_mz", "mat_mz", "p_lapse", "le", "db", "sims"]):
            R[f"{col}{r}"].value = float(res[key])
        print(f"{sc['name']:<32} PD Sr {res['pd_sr']:.4%}  PD Mz {res['pd_mz']:.4%}  "
              f"Mat Sr {res['mat_sr']:.1f}  Mat Mz {res['mat_mz']:.1f}  P(lapse) {res['p_lapse']:.2%}")
    R["B3"].value = f"Corrido: {int(p['n_sims'])} simulaciones, {len(rows)} pólizas, semilla {int(p['seed'])}"
    R[f"C{RES_ROW0 + N_SCEN + 6}"].value = le_w
    print(f"LE ponderada (motor): {le_w:.2f} años")
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def rating_for(pd_, mat, p1):
    ten = min(30, max(1, round(mat)))
    for rt in RATINGS:
        if 1 - (1 - p1[rt]) ** ten >= pd_:
            return rt
    return "ccc o menor"


def size(path, target):
    """Senior/Mezz split grid (total notes fixed): PD, maturity and implied rating per split.

    PD is not monotone in Senior size (tenor and Mezz interest ahead of Senior principal
    both move), so a grid is used instead of a solver.
    """
    wb, p, baf, aaf, fsr_pd, car_fsr, qtab, rows, scen = read_inputs(path)
    D = wb["Matriz_Default"]
    p1 = {D[f"A{5 + i}"].value: D[f"B{5 + i}"].value for i in range(len(RATINGS))}
    if target not in p1:
        sys.exit(f"Calificación objetivo desconocida: {target}. Opciones: {', '.join(RATINGS)}")
    T = int(p["T"])
    q = policy_q(p, baf, aaf, qtab, rows, T)
    total = p["sr_bal"] + p["mz_bal"]
    best = None
    print(f"{'Senior':>8} {'Mezz':>6} {'PD Sr':>8} {'Venc.':>6} {'Rating Sr':>10}")
    for sr_bal in np.arange(10, total + 0.01, 10):
        pp = dict(p, sr_bal=float(sr_bal), mz_bal=float(max(total - sr_bal, 1e-6)))
        res = simulate(pp, q, rows, fsr_pd, car_fsr, scen[0])
        rt = rating_for(res["pd_sr"], res["mat_sr"], p1)
        ok = rt != "ccc o menor" and RATINGS.index(rt) <= RATINGS.index(target)
        if ok:
            best = sr_bal
        print(f"{sr_bal:8.0f} {total - sr_bal:6.0f} {res['pd_sr']:8.3%} {res['mat_sr']:6.1f} {rt:>10}{'  <= objetivo' if ok else ''}")
    print(f"Senior máximo en la grilla con '{target}' o mejor: " + (f"USD {best:.0f} mm" if best else "ninguno"))


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if cmd in ("build", "run") and len(sys.argv) == 3:
        (build if cmd == "build" else run)(sys.argv[2])
    elif cmd == "size" and len(sys.argv) == 4:
        size(sys.argv[2], sys.argv[3])
    else:
        sys.exit(__doc__)
